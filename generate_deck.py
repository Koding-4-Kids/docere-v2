"""Generate Docere investor pitch deck as .pptx — white + orange brand"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Brand colors — white bg + Docere orange (#F15524)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xFA, 0xFA, 0xF7)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
BORDER_GRAY = RGBColor(0xE5, 0xE5, 0xE0)
ORANGE = RGBColor(0xF1, 0x55, 0x24)       # Primary accent #F15524
ORANGE_DARK = RGBColor(0xD9, 0x4A, 0x1C)  # Hover variant
ORANGE_LIGHT = RGBColor(0xFD, 0xE8, 0xE0) # Light orange bg for cards
GREEN = RGBColor(0x10, 0xB9, 0x81)
RED_SOFT = RGBColor(0xDC, 0x26, 0x26)

LOGO_PATH = "/Users/asfawy/docere-v2/frontend/public/docere-logo.png"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_logo(slide, left=Inches(0.6), top=Inches(0.35), height=Inches(0.45)):
    slide.shapes.add_picture(LOGO_PATH, left, top, height=height)


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=BLACK, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_card(slide, left, top, width, height, color=OFF_WHITE, border_color=BORDER_GRAY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_multi_text(slide, left, top, width, height, lines, default_size=18, default_color=BLACK):
    """lines: list of (text, font_size, color, bold, alignment)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line_data in enumerate(lines):
        text = line_data[0]
        size = line_data[1] if len(line_data) > 1 else default_size
        color = line_data[2] if len(line_data) > 2 else default_color
        bold = line_data[3] if len(line_data) > 3 else False
        align = line_data[4] if len(line_data) > 4 else PP_ALIGN.LEFT
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Calibri"
        p.alignment = align
        p.space_after = Pt(6)
    return txBox


def add_divider(slide, left, top, width):
    """Thin horizontal line"""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER_GRAY
    line.line.fill.background()


# ============================================================
# SLIDE 1 — Title
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

# Large centered logo
slide.shapes.add_picture(LOGO_PATH, Inches(4.2), Inches(1.5), height=Inches(1.2))

add_text_box(slide, Inches(1), Inches(3.2), Inches(11), Inches(0.8),
             "The AI memory layer between LMS platforms and classrooms", font_size=28,
             color=DARK, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(1), Inches(4.3), Inches(11), Inches(0.6),
             "Students get personalized learning. Teachers get full visibility.", font_size=20,
             color=ORANGE, alignment=PP_ALIGN.CENTER)

add_divider(slide, Inches(5.5), Inches(5.3), Inches(2.3))

add_text_box(slide, Inches(1), Inches(5.6), Inches(11), Inches(0.5),
             "Youdahe Asfaw  &  Kofi Osei", font_size=18,
             color=GRAY, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(1), Inches(6.1), Inches(11), Inches(0.4),
             "1,000+ students  ·  2 schools deployed  ·  $10K in pitch competition prizes", font_size=14,
             color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2 — Problem
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "The Problem", font_size=44, color=BLACK, bold=True)

add_text_box(slide, Inches(0.8), Inches(2.0), Inches(10), Inches(0.5),
             "Students use AI to cheat. Teachers are blind to it.", font_size=22, color=ORANGE)

problems = [
    ("Students paste homework into ChatGPT",
     "They get the answer, learn nothing, and submit it as their own work."),
    ("Teachers have zero visibility",
     "AI usage is a total black box \u2014 they can\u2019t see what students are doing or learning."),
    ("Current AI tutors never improve",
     "Khanmigo is the same tutor forever. No one measures if tutoring actually works."),
]

for i, (title, desc) in enumerate(problems):
    y = Inches(2.9 + i * 1.4)
    add_card(slide, Inches(0.8), y, Inches(11.5), Inches(1.15))

    # Orange dot
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.15), y + Inches(0.25), Inches(0.15), Inches(0.15))
    dot.fill.solid()
    dot.fill.fore_color.rgb = ORANGE
    dot.line.fill.background()

    add_text_box(slide, Inches(1.5), y + Inches(0.1), Inches(10.5), Inches(0.45),
                 title, font_size=20, color=BLACK, bold=True)
    add_text_box(slide, Inches(1.5), y + Inches(0.55), Inches(10.5), Inches(0.45),
                 desc, font_size=16, color=GRAY)

# ============================================================
# SLIDE 3 — Solution
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "The Solution", font_size=44, color=BLACK, bold=True)

add_text_box(slide, Inches(0.8), Inches(2.0), Inches(10), Inches(0.5),
             "Docere plugs into any LMS in one click. Zero teacher setup.", font_size=22, color=ORANGE)

solutions = [
    ("For Students", "AI tutor grounded in actual course content. Teaches via Socratic questioning \u2014 never gives answers.", ORANGE),
    ("For Teachers", "Full dashboard of every AI-student interaction. See what students struggle with, who\u2019s at risk, and how the class is learning.", GREEN),
    ("The Differentiator", "Every interaction is scored. Scores link to real grades. The system evolves its teaching strategies weekly \u2014 it gets measurably better.", ORANGE_DARK),
]

for i, (title, desc, accent) in enumerate(solutions):
    x = Inches(0.8 + i * 4.0)
    add_card(slide, x, Inches(2.8), Inches(3.7), Inches(4.0))

    # Accent bar at top of card
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.8), Inches(3.7), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    add_text_box(slide, x + Inches(0.3), Inches(3.1), Inches(3.1), Inches(0.5),
                 title, font_size=22, color=accent, bold=True)
    add_text_box(slide, x + Inches(0.3), Inches(3.8), Inches(3.1), Inches(2.8),
                 desc, font_size=16, color=GRAY)

# ============================================================
# SLIDE 4 — How It Works
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "How It Works", font_size=44, color=BLACK, bold=True)

steps = [
    ("1", "LTI Launch", "Teacher clicks \u2018Docere\u2019 inside\nMoodle or Canvas. One click."),
    ("2", "Auto-Sync", "Course materials, assignments,\ngrades, and rosters are pulled\nautomatically. Zero setup."),
    ("3", "Smart Tutoring", "Students get help grounded in\nactual course content. Socratic\nquestioning, not answers."),
    ("4", "Score & Evolve", "Every interaction scored.\nLinked to real grades.\nStrategies evolve weekly."),
]

for i, (num, title, desc) in enumerate(steps):
    x = Inches(0.5 + i * 3.15)
    add_card(slide, x, Inches(2.2), Inches(2.9), Inches(4.5))

    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(1.05), Inches(2.5), Inches(0.7), Inches(0.7))
    circle.fill.solid()
    circle.fill.fore_color.rgb = ORANGE
    circle.line.fill.background()
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(24)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    add_text_box(slide, x + Inches(0.2), Inches(3.4), Inches(2.5), Inches(0.5),
                 title, font_size=20, color=BLACK, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), Inches(4.0), Inches(2.5), Inches(2.5),
                 desc, font_size=15, color=GRAY, alignment=PP_ALIGN.CENTER)

    # Arrow between steps
    if i < 3:
        arrow_x = x + Inches(3.0)
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, Inches(4.2), Inches(0.25), Inches(0.25))
        arr.fill.solid()
        arr.fill.fore_color.rgb = ORANGE
        arr.line.fill.background()

# ============================================================
# SLIDE 5 — Self-Improvement Loop (the moat)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "The Self-Improvement Loop", font_size=44, color=BLACK, bold=True)

add_text_box(slide, Inches(0.8), Inches(2.0), Inches(10), Inches(0.5),
             "No competitor does this. The system gets measurably better every week.", font_size=22, color=ORANGE)

loop_steps = [
    ("Tutor", "5 seed strategies\n(Socratic, Analogy,\nScaffolded, Error-\nFocused, Minimal)"),
    ("Score", "LLM-as-judge +\nheuristics score\nevery interaction\non 4 dimensions"),
    ("Link", "Scores linked to\nreal LMS grades\n(14-day attribution\nwindow)"),
    ("Evolve", "Weekly: mutate top\nstrategies, prune\nworst performers.\nUCB1 bandit selects."),
]

for i, (title, desc) in enumerate(loop_steps):
    x = Inches(0.5 + i * 3.15)
    add_card(slide, x, Inches(2.8), Inches(2.9), Inches(3.5))

    add_text_box(slide, x + Inches(0.3), Inches(3.0), Inches(2.3), Inches(0.5),
                 title, font_size=24, color=ORANGE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.3), Inches(3.6), Inches(2.3), Inches(2.5),
                 desc, font_size=15, color=GRAY, alignment=PP_ALIGN.CENTER)

    if i < 3:
        arrow_x = x + Inches(3.0)
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, Inches(4.3), Inches(0.25), Inches(0.25))
        arr.fill.solid()
        arr.fill.fore_color.rgb = ORANGE
        arr.line.fill.background()

# Bottom summary
add_text_box(slide, Inches(2), Inches(6.5), Inches(9), Inches(0.5),
             "Repeat weekly  \u2192  system improves automatically  \u2192  better student outcomes",
             font_size=16, color=GREEN, bold=True, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 6 — Traction
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "Traction", font_size=44, color=BLACK, bold=True)

metrics = [
    ("1,000+", "Students", "Across 2 middle schools\nand hundreds of educators"),
    ("2", "Schools Live", "Deployed and in active\nuse since 2025"),
    ("$10K", "Prize Money", "Won in pitch competitions\nvalidating the concept"),
    ("0", "Teacher Setup", "LTI integration = one click.\nContent auto-syncs."),
]

for i, (number, label, desc) in enumerate(metrics):
    x = Inches(0.5 + i * 3.15)
    add_card(slide, x, Inches(2.2), Inches(2.9), Inches(3.2))
    add_text_box(slide, x, Inches(2.4), Inches(2.9), Inches(0.9),
                 number, font_size=48, color=ORANGE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x, Inches(3.2), Inches(2.9), Inches(0.5),
                 label, font_size=20, color=BLACK, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.3), Inches(3.8), Inches(2.3), Inches(1.5),
                 desc, font_size=14, color=GRAY, alignment=PP_ALIGN.CENTER)

# Additional traction
add_card(slide, Inches(0.8), Inches(5.8), Inches(11.5), Inches(1.2))
add_multi_text(slide, Inches(1.2), Inches(5.95), Inches(10.5), Inches(1.0), [
    ('Research paper "Learning Agents That Learn" targeting AIED 2026 in Seoul', 16, DARK),
    ("Professor volunteered to help with IRB for formal classroom studies \u2014 he came to us.", 16, GRAY),
])

# ============================================================
# SLIDE 7 — Market
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "Market", font_size=44, color=BLACK, bold=True)

# TAM / SAM / SOM circles
markets = [
    ("$500M", "TAM", "50M US K-12 students @ $10/student/year", Inches(3.6), ORANGE_LIGHT),
    ("$50M", "SAM", "Districts already buying AI/adaptive tools", Inches(2.6), ORANGE_LIGHT),
    ("$500K", "Year 1", "10 paying districts within 12 months", Inches(1.6), ORANGE_LIGHT),
]

x_center = Inches(3.0)
for i, (amount, label, desc, size, color) in enumerate(markets):
    x = x_center + Inches(i * 0.8)
    y_center = Inches(4.2)
    y = y_center - size / 2

    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.color.rgb = ORANGE
    circle.line.width = Pt(2)

# Text overlays
add_text_box(slide, Inches(3.0), Inches(2.8), Inches(3.6), Inches(0.5),
             "$500M TAM", font_size=28, color=ORANGE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(3.0), Inches(3.3), Inches(3.6), Inches(0.4),
             "50M US K-12 students @ $10/student/year", font_size=12, color=DARK, alignment=PP_ALIGN.CENTER)

# Right side — pricing
add_card(slide, Inches(7.2), Inches(1.8), Inches(5.3), Inches(5.2))
add_multi_text(slide, Inches(7.6), Inches(2.0), Inches(4.5), Inches(4.8), [
    ("Pricing", 24, ORANGE, True),
    ("", 8, BLACK),
    ("K-12:  $5\u201315 / student / year", 18, BLACK, True),
    ("A mid-size district (10K students) = $50K\u2013150K/year", 15, GRAY),
    ("", 8, BLACK),
    ("Higher Ed:  $500\u20132,000 / course / semester", 18, BLACK, True),
    ("", 10, BLACK),
    ("Why now?", 24, ORANGE, True),
    ("", 8, BLACK),
    ("Districts already spend $5\u201320/student on IXL, DreamBox.", 15, GRAY),
    ("We\u2019re a direct replacement that actually improves over time.", 15, GRAY),
])

# ============================================================
# SLIDE 8 — Competition
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "Why Not the Others?", font_size=44, color=BLACK, bold=True)

competitors = [
    ("Khanmigo", "Good tutor, same tutor forever.\nNo outcome tracking.\nNo self-improvement.\nRequires switching to Khan.", LIGHT_GRAY),
    ("ChatGPT / Claude", "Memory \u2260 learning.\nRemembering facts is not\nmeasuring if teaching worked.\nZero teacher visibility.", LIGHT_GRAY),
    ("Cognii / Carnegie", "Enterprise, slow, expensive.\nNot LLM-native.\nNo evolution loop.\nMonths-long integration.", LIGHT_GRAY),
    ("Docere", "Scores every interaction.\nLinks to real grades.\nEvolves weekly.\nOne-click LMS integration.\nFull teacher visibility.", ORANGE),
]

for i, (name, desc, accent) in enumerate(competitors):
    x = Inches(0.5 + i * 3.15)
    is_docere = name == "Docere"

    card_bg = ORANGE_LIGHT if is_docere else OFF_WHITE
    border = ORANGE if is_docere else BORDER_GRAY
    add_card(slide, x, Inches(2.2), Inches(2.9), Inches(4.8), color=card_bg, border_color=border)

    # Accent bar at top
    bar_color = ORANGE if is_docere else RGBColor(0xDD, 0xDD, 0xDD)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.2), Inches(2.9), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = bar_color
    bar.line.fill.background()

    name_color = ORANGE if is_docere else BLACK
    add_text_box(slide, x + Inches(0.3), Inches(2.5), Inches(2.3), Inches(0.5),
                 name, font_size=22, color=name_color, bold=True, alignment=PP_ALIGN.CENTER)

    desc_color = DARK if is_docere else GRAY
    add_text_box(slide, x + Inches(0.3), Inches(3.2), Inches(2.3), Inches(3.5),
                 desc, font_size=15, color=desc_color, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 9 — Team
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_logo(slide)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "Team", font_size=44, color=BLACK, bold=True)

# Youdahe card
add_card(slide, Inches(0.8), Inches(2.2), Inches(5.5), Inches(4.8))
add_multi_text(slide, Inches(1.2), Inches(2.4), Inches(4.7), Inches(4.5), [
    ("Youdahe Asfaw", 26, BLACK, True),
    ("Co-Founder  \u00b7  Built everything", 15, ORANGE),
    ("", 8, BLACK),
    ("Built Docere v1 & v2 from scratch \u2014 deployed in 2 schools", 14, GRAY),
    ("ML Research at Hugging Face (open-source, transformers)", 14, GRAY),
    ("Software Engineer at Medtronic (Touch Surgery, AI team)", 14, GRAY),
    ("Software Engineer at Medica (internal workflows)", 14, GRAY),
    ("Incoming SWE Intern at Liberty Mutual (Summer 2026)", 14, GRAY),
    ("Angel investor in AI startups", 14, GRAY),
    ("Code2040 Fellow  \u00b7  ColorStack  \u00b7  Black Minds Tech Fellow", 14, GRAY),
    ("Gustavus Adolphus College \u2014 CS", 14, LIGHT_GRAY),
])

# Kofi card
add_card(slide, Inches(7), Inches(2.2), Inches(5.5), Inches(4.8))
add_multi_text(slide, Inches(7.4), Inches(2.4), Inches(4.7), Inches(4.5), [
    ("Kofi Osei", 26, BLACK, True),
    ("Co-Founder  \u00b7  Research & ML", 15, ORANGE),
    ("", 8, BLACK),
    ("Published ML researcher \u2014 Adaptive Mixture of Experts", 14, GRAY),
    ("SWE at HashiCorp", 14, GRAY),
    ("SWE Intern at Goldman Sachs (Emerging Leaders)", 14, GRAY),
    ("Code2040 Fellow  \u00b7  ColorStack Fellow", 14, GRAY),
    ("Student TA \u2014 1.5 years teaching experience", 14, GRAY),
    ("Co-leads research & experiment design for AIED 2026 paper", 14, GRAY),
    ("Gustavus Adolphus College \u2014 CS", 14, LIGHT_GRAY),
])

# ============================================================
# SLIDE 10 — The Ask
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

# Centered logo
slide.shapes.add_picture(LOGO_PATH, Inches(4.2), Inches(1.0), height=Inches(0.9))

add_text_box(slide, Inches(1), Inches(2.3), Inches(11), Inches(0.8),
             "Let\u2019s Talk", font_size=48, color=BLACK, bold=True, alignment=PP_ALIGN.CENTER)

add_card(slide, Inches(2.8), Inches(3.3), Inches(7.7), Inches(2.8))

add_multi_text(slide, Inches(3.3), Inches(3.5), Inches(6.7), Inches(2.4), [
    ("We have 1,000+ students using the product, a self-improving AI system", 18, DARK, False, PP_ALIGN.CENTER),
    ("that no competitor has, and a research paper going to AIED 2026.", 18, DARK, False, PP_ALIGN.CENTER),
    ("", 12, BLACK),
    ("We\u2019re ready to go full-time and scale to 200 districts.", 20, ORANGE, True, PP_ALIGN.CENTER),
])

add_divider(slide, Inches(5.5), Inches(6.4), Inches(2.3))

add_text_box(slide, Inches(1), Inches(6.6), Inches(11), Inches(0.4),
             "youdahe@docere.ai  \u00b7  kofi@docere.ai", font_size=16,
             color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)


# Save
output_path = "/Users/asfawy/docere-v2/docere-pitch-deck.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
